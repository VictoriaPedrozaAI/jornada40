"""Import a "current schedule" kept as an Excel grid: one row per employee, one column per day,
cells holding shift codes, plus a legend mapping codes to times.

Layout detection (tolerant to templates like the Smartsheet 24/7 one):
- header row = first row whose cell matches HEADER_WORDS (e.g. "Empleado");
- employee names in that column; day columns = the following columns up to a totals column
  ("Horas"/"Hours") or 31 days; rows stop at a blank name or a totals row;
- legend = a (code, description-with-times) table anywhere in the workbook, or passed explicitly.

The file's typed-in totals are ignored (we recompute) but mismatches are reported.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, time, timedelta
from pathlib import Path

import openpyxl

from jornada40.ingest.employees import Issue
from jornada40.shifts import Shift, ShiftDef

HEADER_WORDS = {"empleado", "employee", "nombre", "name", "colaborador"}
TOTAL_WORDS = {"horas", "hours", "total", "totales"}
REST_CODES = {"", "d", "x", "-", "descanso", "off", "r"}
MAX_DAYS = 31
TIME24_RE = re.compile(r"\b([01]?\d|2[0-3]):([0-5]\d)\b")


@dataclass
class GridImport:
    shifts: list[Shift] = field(default_factory=list)
    legend: dict[str, ShiftDef] = field(default_factory=dict)
    employees: list[str] = field(default_factory=list)
    n_days: int = 0
    issues: list[Issue] = field(default_factory=list)


def _clean(v) -> str:
    s = unicodedata.normalize("NFKC", str(v)) if v is not None else ""
    return re.sub(r"\s+", " ", s).strip()


def parse_time_range(text: str) -> tuple[time, time] | None:
    """'Turno nocturno (10:00 p. m. - 6:00 a.m.)' -> (22:00, 06:00). Also '22:00-06:00'."""
    t = _clean(text).replace(" ", "").replace(".", "")
    ampm = [(int(h), int(m or 0), ap.lower()) for h, m, ap in re.findall(r"(\d{1,2})(?::(\d{2}))?([ap])m", t, re.I)]
    if len(ampm) >= 2:
        conv = [time((h % 12) + (12 if ap == "p" else 0), m) for h, m, ap in ampm[:2]]
        return conv[0], conv[1]
    h24 = TIME24_RE.findall(_clean(text))
    if len(h24) >= 2:
        return time(int(h24[0][0]), int(h24[0][1])), time(int(h24[1][0]), int(h24[1][1]))
    return None


def find_legend(wb) -> dict[str, ShiftDef]:
    legend: dict[str, ShiftDef] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            cells = [c for c in row if c.value is not None]
            for a, b in zip(cells, cells[1:]):
                code, desc = _clean(a.value), _clean(b.value)
                if 0 < len(code) <= 3 and code.lower() not in REST_CODES:
                    rng = parse_time_range(desc)
                    if rng and code not in legend:
                        legend[code] = ShiftDef(code, rng[0], rng[1], desc)
    return legend


def import_shift_grid(
    path: str | Path,
    start_date: date,
    sheet: str | None = None,
    legend: dict[str, ShiftDef] | None = None,
    id_map: dict[str, str] | None = None,
) -> GridImport:
    """start_date is required: templates usually carry a placeholder (e.g. 'DD/MM/AA')."""
    res = GridImport()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    res.legend = dict(legend) if legend else find_legend(wb)
    if not res.legend:
        res.issues.append(Issue("ERROR", 0, "legend", "no shift-code legend found; pass legend="))
        return res

    header = None
    for row in ws.iter_rows():
        for c in row:
            if _clean(c.value).lower() in HEADER_WORDS:
                header = c
                break
        if header:
            break
    if header is None:
        res.issues.append(Issue("ERROR", 0, "header", f"no header cell among {sorted(HEADER_WORDS)}"))
        return res

    name_col, hrow = header.column, header.row
    first_day_col = name_col + 1
    total_col = None
    for col in range(first_day_col, first_day_col + MAX_DAYS + 1):
        if _clean(ws.cell(hrow, col).value).lower() in TOTAL_WORDS:
            total_col = col
            break
    last_day_col = (total_col - 1) if total_col else first_day_col + MAX_DAYS - 1
    res.n_days = last_day_col - first_day_col + 1

    patterns: dict[tuple, str] = {}
    r = hrow + 1
    while r <= ws.max_row:
        name = _clean(ws.cell(r, name_col).value)
        if not name or name.lower() in TOTAL_WORDS:
            break
        emp_id = (id_map or {}).get(name, name)
        res.employees.append(emp_id)
        codes = []
        for i, col in enumerate(range(first_day_col, last_day_col + 1)):
            code = _clean(ws.cell(r, col).value)
            codes.append(code)
            if code.lower() in REST_CODES:
                continue
            d = res.legend.get(code) or res.legend.get(code.upper())
            if d is None:
                res.issues.append(Issue("ERROR", r, f"day {i + 1}", f"unknown shift code '{code}' for {emp_id}"))
                continue
            res.shifts.append(Shift.from_def(emp_id, start_date + timedelta(days=i), d))
        if tuple(codes) in patterns:
            res.issues.append(Issue("WARNING", r, "pattern",
                                    f"{emp_id} has an identical schedule to {patterns[tuple(codes)]}"))
        else:
            patterns[tuple(codes)] = emp_id
        if total_col and isinstance(ws.cell(r, total_col).value, (int, float)):
            typed = float(ws.cell(r, total_col).value)
            calc = sum(s.hours for s in res.shifts if s.employee_id == emp_id)
            if abs(typed - calc) > 0.01:
                res.issues.append(Issue("WARNING", r, "total",
                                        f"{emp_id}: file says {typed} h, shifts sum {calc} h"))
        r += 1
    if not res.employees:
        res.issues.append(Issue("ERROR", hrow, "rows", "no employee rows under header"))
    return res


def to_rows(shifts: list[Shift]) -> list[dict]:
    """Canonical long format (one row per shift) — the internal 'current schedule' contract."""
    return [
        {"employee_id": s.employee_id, "date": s.day.isoformat(), "code": s.code,
         "start": s.start.isoformat(timespec="minutes"), "end": s.end.isoformat(timespec="minutes"),
         "hours": s.hours, "jornada_code": s.jornada_code, "jornada_type": s.jornada_type, "night_hours": s.night_hours}
        for s in sorted(shifts, key=lambda s: (s.employee_id, s.start))
    ]
